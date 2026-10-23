# -*- coding: utf-8 -*-
"""Convert a project Markdown doc into a formal hand-over .docx.

Usage:  python md2docx_handover.py [dev|guide]
          dev   (default) -> docs/DEVELOPER_MANUAL_zh-CN.docx
          guide           -> docs/GITHUB_GUIDE_zh-CN.docx

Design notes
- Single source of truth stays the .md file; the docx is a rendering.
- Supports: headings 1-4, tables, fenced code, bullet/numbered/checkbox
  lists, blockquote callouts, inline bold/code/link, --- rules (skipped).
- Numbered lists are rendered with their literal numbers (stable, no
  cross-chapter renumbering like Word's auto 'List Number' style).
- Adds: cover page + info table, TOC field (F9 to refresh), chapter
  page-breaks before H2, header/footer with page numbers, A4 page setup,
  CJK-aware fonts (Microsoft YaHei / Segoe UI).
"""
import datetime
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
TODAY = datetime.date.today().strftime("%Y-%m-%d")

SRC = os.path.join(HERE, "docs", "DEVELOPER_MANUAL_zh-CN.md")
DST = os.path.join(HERE, "docs", "DEVELOPER_MANUAL_zh-CN.docx")
SRC_GUIDE = os.path.join(HERE, "docs", "GITHUB_GUIDE_zh-CN.md")
DST_GUIDE = os.path.join(HERE, "docs", "GITHUB_GUIDE_zh-CN.docx")

EAST = "Microsoft YaHei"
WEST = "Segoe UI"
MONO = "Consolas"
PRIMARY = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x59, 0x59, 0x59)
HDR_FILL = "1F4E79"
CODE_FILL = "F5F5F5"

# Cover / header presets -----------------------------------------------------
DEV_META = {
    "title": "Preprod Telecom Operations Toolkit",
    "subtitle": "开 发 人 员 交 接 文 档",
    "tagline": "（Preprod OCS / CRM 运维工具 · 开发维护手册）",
    "header": "Preprod Telecom Operations Toolkit — 交接文档",
    "footer_note": "— 内部资料，仅限项目团队使用 —",
    "info": [
        ("文档名称", "Preprod Telecom Operations Toolkit 开发人员交接文档"),
        ("文档版本", "v2.0（对应 Markdown 源 docs/DEVELOPER_MANUAL_zh-CN.md）"),
        ("生成日期", TODAY),
        ("适用读者", "接手本工具开发与维护的开发人员"),
        ("环境范围", "预生产环境（Preprod OCS / CRM / Redis / Job Server）"),
        ("维护约定", "先改 Markdown 源，再重新生成本 docx，避免双源不一致"),
        ("问题反馈", "Teams 反馈渠道，由测试人员汇总后每周 1–2 次批量受理"),
    ],
}

GUIDE_META = {
    "title": "YTLC_UAT_TestTool",
    "subtitle": "G i t H u b 协 作 指 南",
    "tagline": "（面向从未使用过 GitHub 的开发者 · 手把手上手手册）",
    "header": "YTLC_UAT_TestTool — GitHub 协作指南",
    "footer_note": "— 内部资料，仅限项目团队使用 —",
    "info": [
        ("文档名称", "YTLC_UAT_TestTool GitHub 协作指南"),
        ("文档版本", "v1.0（对应 Markdown 源 docs/GITHUB_GUIDE_zh-CN.md）"),
        ("生成日期", TODAY),
        ("适用读者", "第一次接触 GitHub / Git 的开发者"),
        ("仓库地址", "https://github.com/dn0218/YTLC_UAT_TestTool"),
        ("维护约定", "先改 Markdown 源，再重新生成本 docx，避免双源不一致"),
        ("问题反馈", "先查指南第 8 节 FAQ，仍未解决则在群里 @ 管理员（dn0218）"),
    ],
}

# ---------------------------------------------------------------- helpers
def _rfonts(element, east=EAST):
    rPr = element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rPr.append(rf)
    rf.set(qn('w:eastAsia'), east)


def style_font(style, size=None, bold=None, color=None, east=EAST,
               west=WEST, mono=False):
    f = style.font
    f.name = MONO if mono else west
    _rfonts(style.element, east=MONO if mono else east)
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color is not None:
        f.color.rgb = color


def run_font(run, size=None, bold=None, italic=None, color=None,
             mono=False):
    f = run.font
    f.name = MONO if mono else WEST
    _rfonts(run._element, east=MONO if mono else EAST)
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color
    return run


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def field_run(paragraph, instr):
    """Insert a Word field (e.g. PAGE / TOC) into a paragraph."""
    r1 = paragraph.add_run()
    fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'begin')
    r1._element.append(fc)
    r2 = paragraph.add_run()
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve')
    it.text = instr
    r2._element.append(it)
    r3 = paragraph.add_run()
    fc2 = OxmlElement('w:fldChar'); fc2.set(qn('w:fldCharType'), 'end')
    r3._element.append(fc2)
    return r1, r2, r3


INLINE_RE = re.compile(
    r'(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')

def add_inline(par, text, base_size=None, base_color=None, bold_all=False):
    for tok in INLINE_RE.split(text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            run_font(par.add_run(tok[2:-2]), size=base_size, bold=True,
                     color=base_color)
        elif tok.startswith('`') and tok.endswith('`') and len(tok) > 2:
            run_font(par.add_run(tok[1:-1]), size=(base_size or 10.5) - 1.0,
                     color=RGBColor(0xC0, 0x39, 0x2B), mono=True,
                     bold=bold_all)
        elif tok.startswith('['):
            m = re.match(r'\[([^\]]+)\]\([^)]*\)', tok)
            run_font(par.add_run(m.group(1) if m else tok),
                     size=base_size, bold=bold_all, color=base_color)
        else:
            run_font(par.add_run(tok), size=base_size, bold=bold_all,
                     color=base_color)


def add_code_block(doc, lines):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = doc.styles['Table Grid']
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    shade_cell(cell, CODE_FILL)
    cell.paragraphs[0].text = ''
    for i, ln in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.12
        run_font(p.add_run(ln if ln else ' '), size=8.5, mono=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, rows):
    ncols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=0, cols=ncols)
    tbl.style = doc.styles['Table Grid']
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True
    for ri, row in enumerate(rows):
        cells = tbl.add_row().cells
        for ci in range(ncols):
            txt = row[ci] if ci < len(row) else ''
            p = cells[ci].paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            if ri == 0:
                shade_cell(cells[ci], HDR_FILL)
                add_inline(p, txt, base_size=9,
                           base_color=RGBColor(0xFF, 0xFF, 0xFF),
                           bold_all=True)
            else:
                add_inline(p, txt, base_size=9)
                if ri % 2 == 0:
                    shade_cell(cells[ci], "F2F6FA")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def split_row(line):
    line = line.strip()
    if line.startswith('|'):
        line = line[1:]
    if line.endswith('|'):
        line = line[:-1]
    return [c.strip() for c in line.split('|')]


def is_sep_row(line):
    return bool(re.match(r'^\s*\|?[\s:\-|]+\|?\s*$', line)) and '-' in line


# ---------------------------------------------------------------- document
def build_styles(doc):
    n = doc.styles['Normal']
    style_font(n, size=10.5)
    n.paragraph_format.line_spacing = 1.4
    n.paragraph_format.space_after = Pt(4)

    specs = {'Heading 1': (18, PRIMARY), 'Heading 2': (14, PRIMARY),
             'Heading 3': (12, RGBColor(0x2E, 0x5F, 0x8A)),
             'Heading 4': (11, MUTED)}
    for name, (size, color) in specs.items():
        st = doc.styles[name]
        style_font(st, size=size, bold=True, color=color)
        st.paragraph_format.space_before = Pt(14 if name != 'Heading 1' else 6)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True


def setup_page(doc, header_text):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.different_first_page_header_footer = True

    hp = sec.header.paragraphs[0]
    hp.text = ''
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_font(hp.add_run(header_text), size=8, color=MUTED)

    fp = sec.footer.paragraphs[0]
    fp.text = ''
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_font(fp.add_run("第 "), size=8.5, color=MUTED)
    for r in field_run(fp, ' PAGE '):
        pass
    run_font(fp.add_run(" 页 / 共 "), size=8.5, color=MUTED)
    field_run(fp, ' NUMPAGES ')
    run_font(fp.add_run(" 页"), size=8.5, color=MUTED)
    for r in fp.runs:
        run_font(r, size=8.5, color=MUTED)


def add_cover(doc, meta):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_font(p.add_run(meta["title"]), size=24, bold=True, color=PRIMARY)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_font(p.add_run(meta["subtitle"]), size=17, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_font(p.add_run(meta["tagline"]), size=11, color=MUTED)
    doc.add_paragraph(); doc.add_paragraph()

    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = doc.styles['Table Grid']
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in meta["info"]:
        cells = tbl.add_row().cells
        cells[0].width, cells[1].width = Cm(3.4), Cm(11.8)
        shade_cell(cells[0], "DEEAF6")
        p = cells[0].paragraphs[0]
        run_font(p.add_run(k), size=10, bold=True, color=PRIMARY)
        add_inline(cells[1].paragraphs[0], v, base_size=10)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_font(p.add_run(meta["footer_note"]), size=9, color=MUTED)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_toc(doc):
    doc.add_heading("目  录", level=1)
    p = doc.add_paragraph()
    field_run(p, r'TOC \o "1-2" \h \z \u')
    tip = doc.add_paragraph()
    run_font(tip.add_run("（在 Word 中按 Ctrl+A → F9 可刷新目录页码）"),
             size=8.5, italic=True, color=MUTED)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def convert(src=SRC, dst=DST, meta=None):
    meta = meta or DEV_META
    with open(src, encoding='utf-8') as f:
        lines = f.read().splitlines()

    doc = Document()
    build_styles(doc)
    setup_page(doc, meta["header"])
    add_cover(doc, meta)
    add_toc(doc)

    stats = {'h2': 0, 'tables': 0, 'code': 0, 'skipped_hr': 0}
    i, n = 0, len(lines)
    first_h2_done = False
    # skip the original H1 + callout (cover replaces them)
    while i < n and (lines[i].startswith('# ') or lines[i].startswith('>')
                     or lines[i].strip() in ('', '---')):
        i += 1

    skip_md_toc = False
    while i < n:
        line = lines[i]
        s = line.strip()

        # skip the markdown TOC chapter ("## 目录" until next heading):
        # the docx already has a real Word TOC field on its own page.
        if re.match(r'^##\s+目录\s*$', s):
            skip_md_toc = True
            i += 1; continue
        if skip_md_toc:
            if re.match(r'^#{1,4}\s+', s):
                skip_md_toc = False
            else:
                i += 1; continue

        if not s:
            i += 1; continue
        if s == '---':
            stats['skipped_hr'] += 1; i += 1; continue
        if s.startswith('```'):
            block = []
            i += 1
            while i < n and not lines[i].strip().startswith('```'):
                block.append(lines[i]); i += 1
            i += 1
            add_code_block(doc, block); stats['code'] += 1
            continue
        if s.startswith('|') and i + 1 < n and is_sep_row(lines[i + 1]):
            rows = [split_row(s)]
            i += 2
            while i < n and lines[i].strip().startswith('|'):
                rows.append(split_row(lines[i].strip())); i += 1
            add_table(doc, rows); stats['tables'] += 1
            continue
        m = re.match(r'^(#{1,4})\s+(.*)$', s)
        if m:
            lvl, txt = len(m.group(1)), m.group(2).strip()
            if lvl == 2:
                if first_h2_done:
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                first_h2_done = True
                stats['h2'] += 1
            h = doc.add_heading('', level=lvl)
            add_inline(h, txt)
            i += 1; continue
        if s.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            add_inline(p, s[2:].strip(), base_color=MUTED)
            for r in p.runs:
                r.italic = True
            i += 1; continue
        m = re.match(r'^\s*-\s+\[ \]\s+(.*)$', line)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            run_font(p.add_run("☐ "), size=10.5)
            add_inline(p, m.group(1))
            i += 1; continue
        m = re.match(r'^(\s*)[-*]\s+(.*)$', line)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            if len(m.group(1)) >= 2:
                p.paragraph_format.left_indent = Cm(1.3)
            add_inline(p, m.group(2))
            i += 1; continue
        # numbered list: keep the literal number so numbering never
        # re-numbers across chapters (Word 'List Number' would).
        m = re.match(r'^(\s*)(\d+)\.\s+(.*)$', line)
        if m:
            p = doc.add_paragraph()
            if len(m.group(1)) >= 2:
                p.paragraph_format.left_indent = Cm(1.3)
            run_font(p.add_run(m.group(2) + ". "), size=10.5, bold=True,
                     color=PRIMARY)
            add_inline(p, m.group(3))
            i += 1; continue
        p = doc.add_paragraph()
        add_inline(p, s)
        i += 1

    doc.save(dst)
    return stats


if __name__ == '__main__':
    preset = sys.argv[1].lower() if len(sys.argv) > 1 else 'dev'
    if preset == 'guide':
        out = DST_GUIDE
        st = convert(SRC_GUIDE, DST_GUIDE, GUIDE_META)
    else:
        out = DST
        st = convert()
    print(f"OK -> {out}")
    print(f"H2 chapters={st['h2']}, tables={st['tables']}, "
          f"code blocks={st['code']}, hr skipped={st['skipped_hr']}, "
          f"size={os.path.getsize(out)} bytes")
